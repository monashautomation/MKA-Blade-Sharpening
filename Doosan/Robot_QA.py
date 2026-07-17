#!/usr/bin/env python3
"""
Robot_QA.py - summarise blade-sharpening grind logs pasted into Google Sheets.

Workflow this supports:
  * You clear the Dart Studio log, run a job, then copy the whole log and paste
    it into a new tab of the shared Google Sheet (tab named by date / number).
  * Column A = time (HH:MM:SS.ffffff), column B = the tp_log message.
  * Each tab is a *session* that may contain several grind runs (each run starts
    at a "Bay N Grinder N Angle A Depth D" config line).

What it does:
  1. Asks which tab (sheet) to summarise - lists them with quick stats.
  2. Extracts + parses the tp_log lines from that tab.
  3. Segments the tab into runs, summarises each run + a tab-level rollup:
       - config + grind variant (v1 pass / v2 plunge-sweep)
       - teeth-ground count, pitch (X) and measured-depth (Y) stats
       - depth-tracking accuracy (commanded vs actual cut depth, from DIAG lines)
       - quality events (outliers, aborts, 0-depth, no-blade...) + timing
  4. Prints console tables and writes a Markdown report + per-tooth CSV.

Access:
  The sheet must be link-shared ("Anyone with the link -> Viewer"). The script
  downloads it as XLSX over plain HTTPS - no Google login / credentials needed.
  Use --file to read a locally downloaded .xlsx instead (fully offline).

Requires: requests, openpyxl (tabulate optional - falls back to a plain table).

Usage:
  python Robot_QA.py                 # interactive: list tabs, pick one
  python Robot_QA.py 107             # summarise tab "107"
  python Robot_QA.py --all           # summarise every tab
  python Robot_QA.py --file book.xlsx 107     # offline, from a downloaded file
  python Robot_QA.py --no-save       # console only, don't write report files
"""

import argparse
import csv
import io
import math
import re
import statistics as st
import sys
import warnings
from dataclasses import dataclass, field
from datetime import datetime, time as dtime
from pathlib import Path

warnings.filterwarnings("ignore")  # silence requests' urllib3 version warning

try:
    from tabulate import tabulate
except Exception:                                    # noqa: BLE001
    tabulate = None

# ------------------------------------------------------------------ config ---
SPREADSHEET_ID = "1pHZQtr7gV0cmwPIS-oqQCF-xVPOLSfUKtBjUxC0TKV8"
REPORT_DIR = Path(__file__).resolve().parent / "QA_reports"
IDLE_GAP_S = 120.0    # gaps longer than this (robot idle / between jobs) don't count as grind time

# ---------------------------------------------------------------- patterns ---
FLOAT = r"[-+]?\d*\.?\d+"
RE_CONFIG    = re.compile(rf"Bay\s+(\d+)\s+Grinder\s+(\d+)\s+Angle\s+({FLOAT})\s+Depth\s+({FLOAT})")
RE_DETECT    = re.compile(rf"X:\s*({FLOAT})\s+Y:\s*({FLOAT})\s+Depth:\s*({FLOAT})\s+DetD:\s*({FLOAT})")
RE_CALCCUT   = re.compile(rf"Calc cut:\s*({FLOAT})\s+Passes:\s*(\d+)")
RE_CMDACT    = re.compile(rf"cmd_depth=({FLOAT})\s+actual_depth=({FLOAT})")
RE_OFFLIMIT  = re.compile(rf"Offset:\s*({FLOAT})\s+Limit:\s*({FLOAT})")
RE_OUTLIER   = re.compile(r"ERROR outlier #(\d+)")
RE_BLADERT   = re.compile(rf"Blade right:\s*({FLOAT})")


# --------------------------------------------------------------- log model ---
@dataclass
class Tooth:
    """One detection cycle (grinder sees a tooth, then cuts it)."""
    idx: int
    t: object                       # datetime.time of the detection line
    x: float                        # pitch  (detection X)
    y: float                        # measured tooth depth (detection Y)
    target: float                   # commanded cut depth (detection "Depth")
    detd: float                     # detection depth (DetD)
    status: str = "cut"             # cut | rejected | first
    passes: int = None
    calc_cut: float = None
    diag: list = field(default_factory=list)   # (cmd_depth, actual_depth) samples

    @property
    def final_cmd(self):
        return self.diag[-1][0] if self.diag else None

    @property
    def final_actual(self):
        return self.diag[-1][1] if self.diag else None

    @property
    def errs(self):
        return [a - c for c, a in self.diag]

    @property
    def err_maxabs(self):
        e = self.errs
        return max(abs(v) for v in e) if e else None


@dataclass
class Run:
    idx: int
    config: dict = None             # {bay, grinder, angle, target} or None
    rows: list = field(default_factory=list)   # (time, msg)
    teeth: list = field(default_factory=list)
    n_outliers: int = 0
    n_aborts: int = 0
    n_zero_depth: int = 0
    n_blades: int = 0               # "Grind Ready" count
    n_other: int = 0                # messages we didn't specifically classify
    events: set = field(default_factory=set)   # terminal / notable markers
    blade_right: float = None
    manual_angle: float = None      # bevel angle supplied by the user when no config line

    @property
    def cut_teeth(self):
        return [t for t in self.teeth if t.status in ("cut", "first")]

    @property
    def n_rejected(self):
        return sum(1 for t in self.teeth if t.status == "rejected")


# --------------------------------------------------------------- utilities ---
def to_seconds(v):
    """Column-A cell -> seconds since midnight (accepts time/datetime/str)."""
    if isinstance(v, dtime):
        return v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    if isinstance(v, datetime):
        return v.hour * 3600 + v.minute * 60 + v.second + v.microsecond / 1e6
    if isinstance(v, str):
        m = re.match(r"(\d{1,2}):(\d{2}):(\d{2})(?:\.(\d+))?", v.strip())
        if m:
            h, mn, s = int(m[1]), int(m[2]), int(m[3])
            frac = float("0." + m[4]) if m[4] else 0.0
            return h * 3600 + mn * 60 + s + frac
    return None


def hms(v):
    if isinstance(v, (dtime, datetime)):
        return v.strftime("%H:%M:%S")
    return str(v) if v is not None else "-"


def duration_s(rows):
    """Active grind time = sum of forward deltas, ignoring gaps > IDLE_GAP_S.
    Robust to a tab holding several concatenated captures (time resets) and to
    long idle waits between teeth/jobs, which would otherwise dominate cadence."""
    ts = [to_seconds(t) for t, _ in rows if to_seconds(t) is not None]
    total = 0.0
    for a, b in zip(ts, ts[1:]):
        d = b - a
        if 0 <= d <= IDLE_GAP_S:
            total += d
    return total


def n_captures(rows):
    """How many separate log captures were pasted (1 + number of time resets)."""
    ts = [to_seconds(t) for t, _ in rows if to_seconds(t) is not None]
    return (1 + sum(1 for a, b in zip(ts, ts[1:]) if b < a)) if ts else 0


def fmt(v, nd=3):
    if v is None:
        return "-"
    if isinstance(v, float):
        return f"{v:.{nd}f}"
    return str(v)


def stat_block(vals):
    """(mean, min, max) or (None, None, None) for an empty list."""
    vals = [v for v in vals if v is not None]
    if not vals:
        return None, None, None
    return st.mean(vals), min(vals), max(vals)


def tooth_length(pitch, angle_deg):
    """Actual (sloped) tooth-facet length from the measured pitch and the bevel
    angle. Angle is measured from vertical, so the facet sits (90-angle) from the
    horizontal and the pitch is its horizontal projection:  L = pitch / cos(90-angle)
    ( == pitch / sin(angle) )."""
    if pitch is None or angle_deg is None:
        return None
    denom = math.cos(math.radians(90 - angle_deg))
    if abs(denom) < 1e-6:                      # facet ~vertical -> undefined
        return None
    return pitch / denom


def fmt_range(lo, hi, nd=2):
    """'lo-hi' for a min/max pair, or '-' if empty."""
    if lo is None or hi is None:
        return "-"
    return f"{lo:.{nd}f}-{hi:.{nd}f}"


def render_table(headers, rows):
    if tabulate:
        return tabulate(rows, headers=headers, tablefmt="simple", floatfmt=".3f")
    # minimal fallback
    cols = [headers] + [[fmt(c) for c in r] for r in rows]
    widths = [max(len(str(row[i])) for row in cols) for i in range(len(headers))]
    line = lambda r: "  ".join(str(c).ljust(widths[i]) for i, c in enumerate(r))
    out = [line(headers), line(["-" * w for w in widths])]
    out += [line([fmt(c) for c in r]) for r in rows]
    return "\n".join(out)


# ---------------------------------------------------------------- fetching ---
def download_workbook(spreadsheet_id):
    import requests
    url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=xlsx"
    r = requests.get(url, timeout=30, allow_redirects=True)
    ctype = r.headers.get("content-type", "")
    if r.status_code != 200 or "spreadsheetml" not in ctype:
        raise SystemExit(
            f"Could not download the sheet (HTTP {r.status_code}, type '{ctype}').\n"
            "Make sure it is shared 'Anyone with the link -> Viewer', or pass --file "
            "with a locally downloaded .xlsx."
        )
    return io.BytesIO(r.content)


def load_tabs(source):
    """source: BytesIO or path -> dict {tab_name: [(time, msg), ...]}."""
    import openpyxl
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    tabs = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for r in ws.iter_rows(values_only=True):
            t = r[0] if len(r) > 0 else None
            msg = r[1] if len(r) > 1 and r[1] is not None else None
            if msg is None and t is not None and isinstance(t, str):
                msg, t = t, None            # message-only paste (no time column)
            if msg is not None:
                rows.append((t, str(msg).strip()))
        tabs[name] = rows
    return tabs


# ----------------------------------------------------------------- parsing ---
def segment_runs(rows):
    """Split a tab's rows into runs. A new run begins at each config line;
    rows before the first config form an initial (unknown-config) run."""
    runs, cur = [], Run(idx=0)
    for t, msg in rows:
        m = RE_CONFIG.search(msg)
        if m:
            if cur.rows:                     # close the run in progress
                runs.append(cur)
                cur = Run(idx=len(runs))
            cur.config = {"bay": int(m[1]), "grinder": int(m[2]),
                          "angle": float(m[3]), "target": float(m[4])}
        cur.rows.append((t, msg))
    if cur.rows:
        runs.append(cur)
    for i, r in enumerate(runs):             # renumber contiguously
        r.idx = i + 1
        analyse_run(r)
    return runs


def analyse_run(run):
    """Group a run's rows into teeth and tally events."""
    cur = None
    for t, msg in run.rows:
        md = RE_DETECT.search(msg)
        if md:
            if cur:
                run.teeth.append(cur)
            cur = Tooth(idx=len(run.teeth) + 1, t=t,
                        x=float(md[1]), y=float(md[2]),
                        target=float(md[3]), detd=float(md[4]))
            continue

        if cur is not None:                  # attach to the current tooth block
            if "First tooth" in msg:
                cur.status = "first"
            mo = RE_OUTLIER.search(msg)
            if mo:
                cur.status = "rejected"
                run.n_outliers += 1
            mc = RE_CALCCUT.search(msg)
            if mc:
                cur.calc_cut, cur.passes = float(mc[1]), int(mc[2])
            ma = RE_CMDACT.search(msg)
            if ma:
                cur.diag.append((float(ma[1]), float(ma[2])))

        # run-level tallies (independent of tooth grouping)
        if "bad detections in a row" in msg:
            run.n_aborts += 1
            run.events.add("abort")
        if "0 depth cut" in msg:
            run.n_zero_depth += 1
        if "Grind Ready" in msg:
            run.n_blades += 1
        if "Finished grinding" in msg:
            run.events.add("finished")
        if "Stopping operation" in msg:
            run.events.add("stopped")
        if "Emergency" in msg:
            run.events.add("emergency")
        if "No blade detected" in msg:
            run.events.add("no_blade")
        if "Incorrect Grinder ID" in msg or "Incorrect Bay ID" in msg:
            run.events.add("bad_id")
        if "Invalid direction" in msg:
            run.events.add("invalid_dir")
        mb = RE_BLADERT.search(msg)
        if mb:
            run.blade_right = float(mb[1])

    if cur:
        run.teeth.append(cur)

    # a rejected block that never turned into a cut stays "rejected"; anything
    # with cut diagnostics but still flagged rejected was actually re-cut later.
    for tth in run.teeth:
        if tth.status == "rejected" and tth.diag:
            tth.status = "cut"


def grind_variant(run):
    joined = "\n".join(m for _, m in run.rows)
    if "DIAG plunge" in joined or "DIAG sweep" in joined:
        return "v2 (plunge/sweep)"
    if "DIAG pass" in joined:
        return "v1 (pass)"
    return "unknown"


def run_end_status(run, is_last):
    """Explicit terminators win. Otherwise, a run followed by another run
    completed (the robot returned to START-wait and a new job began); the final
    run without a terminator may just be where the log copy stopped.
    (These logs have no 'Finished grinding' marker, so completion is inferred.)"""
    for tag in ("finished", "emergency", "abort", "stopped", "no_blade",
                "bad_id", "invalid_dir"):
        if tag in run.events:
            return {"finished": "finished", "emergency": "EMERGENCY",
                    "abort": "aborted (3-bad)", "stopped": "stopped",
                    "no_blade": "no blade", "bad_id": "bad ID",
                    "invalid_dir": "invalid dir"}[tag]
    if run.cut_teeth:
        return "ends tab" if is_last else "done*"
    return "empty"


# --------------------------------------------------------------- summarise ---
def summarise_run(run, is_last=False):
    """Flat dict of the numbers we report for one run."""
    teeth = run.cut_teeth
    pit_mean, pit_min, pit_max = stat_block([t.x for t in teeth])
    dep_mean, dep_min, dep_max = stat_block([t.y for t in teeth])
    detd_mean, _, _ = stat_block([t.detd for t in teeth])
    all_err = [e for t in teeth for e in t.errs]
    achieved = [t.final_actual - t.target for t in teeth
                if t.final_actual is not None and t.target is not None]
    cfg = run.config or {}
    angle_used = cfg.get("angle") if run.config is not None else run.manual_angle
    angle_source = ("config" if run.config is not None
                    else "manual" if run.manual_angle is not None else None)
    tlens = [tooth_length(t.x, angle_used) for t in teeth]
    tlens = [v for v in tlens if v is not None]
    tlen_min, tlen_max = (min(tlens), max(tlens)) if tlens else (None, None)
    return {
        "run": run.idx,
        "bay": cfg.get("bay"),
        "grinder": cfg.get("grinder"),
        "angle": cfg.get("angle"),
        "angle_used": angle_used,
        "angle_source": angle_source,
        "tooth_len_min": tlen_min,
        "tooth_len_max": tlen_max,
        "target_depth": cfg.get("target"),
        "variant": grind_variant(run),
        "blades": run.n_blades,
        "teeth": len(teeth),
        "rejected": run.n_rejected,
        "aborts": run.n_aborts,
        "zero_depth": run.n_zero_depth,
        "pitch_mean": pit_mean, "pitch_min": pit_min, "pitch_max": pit_max,
        "depthY_mean": dep_mean, "depthY_min": dep_min, "depthY_max": dep_max,
        "detd_mean": detd_mean,
        "track_err_mean": (st.mean(all_err) if all_err else None),
        "track_err_maxabs": (max(abs(e) for e in all_err) if all_err else None),
        "err_pos_max": (max(all_err) if all_err else None),   # largest overshoot (+)
        "err_neg_max": (min(all_err) if all_err else None),   # largest undershoot (-)
        "achieved_vs_target_mean": (st.mean(achieved) if achieved else None),
        "duration_s": duration_s(run.rows),
        "sec_per_tooth": (duration_s(run.rows) / len(teeth) if teeth else None),
        "start": hms(run.rows[0][0]) if run.rows else "-",
        "end": hms(run.rows[-1][0]) if run.rows else "-",
        "status": run_end_status(run, is_last),
        "captures": n_captures(run.rows),
        "config_known": run.config is not None,
    }


def config_label(s):
    if not s["config_known"]:
        if s.get("angle_used") is not None:
            return f"(no config, {fmt(s['angle_used'], 1)}deg*)"
        return "(no config)"
    return f"B{s['bay']} G{s['grinder']} {fmt(s['angle'],1)}deg d{fmt(s['target_depth'],2)}"


# ----------------------------------------------------------------- reports ---
def print_report(tab_name, runs):
    summaries = [summarise_run(r, r is runs[-1]) for r in runs]
    all_rows = [rw for r in runs for rw in r.rows]
    total_dur = duration_s(all_rows)
    caps = n_captures(all_rows)
    order = f"{hms(runs[0].rows[0][0])} -> {hms(runs[-1].rows[-1][0])}"
    teeth_tot = sum(s["teeth"] for s in summaries)
    rej_tot = sum(s["rejected"] for s in summaries)
    ab_tot = sum(s["aborts"] for s in summaries)
    all_err = [e for r in runs for t in r.cut_teeth for e in t.errs]

    print("\n" + "=" * 78)
    print(f"TAB {tab_name!r}  |  log order {order}  |  active grind {total_dur/60:.1f} min  |  "
          f"{len(runs)} run(s), {teeth_tot} teeth cut")
    if caps > 1:
        print(f"NOTE: {caps} separate log captures pasted into this tab (time resets detected).")
    if all_err:
        print(f"Depth tracking (actual - commanded): mean {st.mean(all_err):+.4f} mm, "
              f"max |err| {max(abs(e) for e in all_err):.4f} mm  "
              f"({len(all_err)} samples)")
    tl_all = [tooth_length(t.x, s["angle_used"])
              for r, s in zip(runs, summaries) for t in r.cut_teeth]
    tl_all = [v for v in tl_all if v is not None]
    if tl_all:
        print(f"Tooth length (pitch / cos(90-angle)): {min(tl_all):.2f} - {max(tl_all):.2f} mm")
    print(f"Rejected detections: {rej_tot}   Aborts: {ab_tot}")
    print("=" * 78)

    headers = ["Run", "Config", "Var", "Teeth", "Rej", "Ab",
               "Av. Pitch", "Av. Depth", "Av. Error", "+ max", "- max", "s/tooth",
               "Tooth len"]
    table = []
    for s in summaries:
        table.append([
            s["run"], config_label(s), s["variant"].split()[0],
            s["teeth"], s["rejected"], s["aborts"],
            fmt(s["pitch_mean"], 2), fmt(s["depthY_mean"], 2),
            fmt(s["track_err_mean"], 4), fmt(s["err_pos_max"], 4),
            fmt(s["err_neg_max"], 4), fmt(s["sec_per_tooth"], 1),
            fmt_range(s["tooth_len_min"], s["tooth_len_max"]),
        ])
    print(render_table(headers, table))
    print("(Av. Pitch / Av. Depth = mean tooth pitch (X) / measured depth (Y), mm.")
    print(" Av. Error = mean actual-commanded cut depth; + max / - max = largest "
          "positive / negative tracking error, mm.  s/tooth excludes idle >2min.")
    print(" Tooth len = min-max of pitch / cos(90-angle) across teeth, mm;  "
          "* = angle entered manually (no config line).)")
    return summaries


def write_markdown(path, tab_name, runs, summaries):
    L = []
    all_rows = [rw for r in runs for rw in r.rows]
    total_dur = duration_s(all_rows)
    caps = n_captures(all_rows)
    order = f"{hms(runs[0].rows[0][0])} -> {hms(runs[-1].rows[-1][0])}"
    teeth_tot = sum(s["teeth"] for s in summaries)
    all_err = [e for r in runs for t in r.cut_teeth for e in t.errs]

    L.append(f"# QA report - tab `{tab_name}`\n")
    L.append(f"- Generated: {datetime.now():%Y-%m-%d %H:%M:%S}")
    L.append(f"- Log order: {order}   Active grind time: {total_dur/60:.1f} min (idle >2min excluded)")
    L.append(f"- Runs: {len(runs)}   Teeth cut: {teeth_tot}")
    if caps > 1:
        L.append(f"- **Note:** {caps} separate log captures were pasted into this tab "
                 "(time resets detected).")
    if all_err:
        L.append(f"- Depth tracking (actual-commanded): mean **{st.mean(all_err):+.4f} mm**, "
                 f"max |err| **{max(abs(e) for e in all_err):.4f} mm** ({len(all_err)} samples)")
    tl_all = [tooth_length(t.x, s["angle_used"])
              for r, s in zip(runs, summaries) for t in r.cut_teeth]
    tl_all = [v for v in tl_all if v is not None]
    if tl_all:
        L.append(f"- Tooth length (pitch / cos(90-angle)): **{min(tl_all):.2f} - "
                 f"{max(tl_all):.2f} mm**")
    L.append("")

    L.append("## Runs\n")
    headers = ["Run", "Config", "Variant", "Blades", "Teeth", "Rej", "Aborts",
               "Av.Pitch", "Av.Depth", "Av.Error", "+max", "-max", "Tooth len",
               "Achv-Tgt", "Dur(s)", "s/tooth", "Status"]
    table = [[s["run"], config_label(s), s["variant"], s["blades"], s["teeth"],
              s["rejected"], s["aborts"], fmt(s["pitch_mean"], 2),
              fmt(s["depthY_mean"], 2), fmt(s["track_err_mean"], 4),
              fmt(s["err_pos_max"], 4), fmt(s["err_neg_max"], 4),
              fmt_range(s["tooth_len_min"], s["tooth_len_max"]),
              fmt(s["achieved_vs_target_mean"], 4),
              fmt(s["duration_s"], 1), fmt(s["sec_per_tooth"], 1), s["status"]]
             for s in summaries]
    L.append(md_table(headers, table))
    L.append("")
    L.append("_Av.Pitch/Av.Depth = mean tooth pitch (X) / measured depth (Y), mm. "
             "Av.Error = mean actual-commanded cut depth; +max / -max = largest "
             "positive / negative tracking error, mm. Tooth len = min-max of "
             "pitch / cos(90-angle) across teeth, mm (angle from vertical; "
             "config `deg*` = entered manually). Achv-Tgt = final actual depth - "
             "target, mm. Dur & s/tooth exclude idle >2min. done\\* = completed "
             "(inferred from a following run); ends tab = last run, may be truncated. "
             "Per-tooth MaxAbsErr = worst |actual-commanded| among that tooth's moves._\n")

    for run in runs:
        s = summarise_run(run, run is runs[-1])
        L.append(f"### Run {run.idx} - {config_label(s)}  ({s['status']})\n")
        note = []
        if not s["config_known"]:
            note.append("no config line in this run (unknown Bay/Grinder/Angle/Depth)")
        if run.idx == 1 and not s["config_known"] and any(
                r.config for r in runs):
            note.append("looks like a leftover fragment before the first job")
        if note:
            L.append("> Note: " + "; ".join(note) + "\n")
        if not run.cut_teeth:
            L.append("_No teeth cut in this run._\n")
            continue
        th = ["#", "Time", "Status", "Pitch(X)", "ToothLen", "DepthY", "Target",
              "DetD", "Passes", "CalcCut", "FinalCmd", "FinalActual", "MaxAbsErr"]
        trows = [[t.idx, hms(t.t), t.status, fmt(t.x, 2),
                  fmt(tooth_length(t.x, s["angle_used"]), 3), fmt(t.y, 2),
                  fmt(t.target, 2), fmt(t.detd, 2), t.passes, fmt(t.calc_cut, 3),
                  fmt(t.final_cmd, 4), fmt(t.final_actual, 4), fmt(t.err_maxabs, 4)]
                 for t in run.teeth]
        L.append(md_table(th, trows))
        L.append("")

    path.write_text("\n".join(L), encoding="utf-8")


def md_table(headers, rows):
    """A real, aligned GitHub-flavoured Markdown table (renders as a table, and
    the raw text is column-aligned too)."""
    frows = [[fmt(c) for c in r] for r in rows]
    if tabulate:
        return tabulate(frows, headers=headers, tablefmt="github", disable_numparse=True)
    out = ["| " + " | ".join(str(h) for h in headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    for r in frows:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def write_csv(path, tab_name, runs):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["tab", "run", "config", "angle_deg", "tooth", "time", "status",
                    "pitch_x", "tooth_len", "depth_y", "target_depth", "detd",
                    "passes", "calc_cut", "final_cmd_depth", "final_actual_depth",
                    "track_err_maxabs", "n_diag"])
        for run in runs:
            s = summarise_run(run)
            cfg = config_label(s)
            ang = s["angle_used"]
            for t in run.teeth:
                w.writerow([tab_name, run.idx, cfg, fmt(ang, 2), t.idx, hms(t.t),
                            t.status, fmt(t.x, 3), fmt(tooth_length(t.x, ang), 4),
                            fmt(t.y, 3), fmt(t.target, 3), fmt(t.detd, 3), t.passes,
                            fmt(t.calc_cut, 4), fmt(t.final_cmd, 4),
                            fmt(t.final_actual, 4), fmt(t.err_maxabs, 4), len(t.diag)])


# -------------------------------------------------------------------- flow ---
def choose_tab(tabs):
    names = list(tabs)
    print("\nAvailable tabs (each = one pasted job/session):\n")
    rows = []
    for i, name in enumerate(names, 1):
        r = tabs[name]
        span = f"{hms(r[0][0])}-{hms(r[-1][0])}" if r else "-"
        configs = sum(1 for _, m in r if RE_CONFIG.search(m))
        rows.append([i, name, len(r), configs, span])
    print(render_table(["#", "Tab", "Rows", "Runs", "Time span"], rows))
    while True:
        sel = input("\nPick a tab (number or name, 'all' for every tab): ").strip()
        if sel.lower() == "all":
            return names
        if sel.isdigit() and 1 <= int(sel) <= len(names):
            return [names[int(sel) - 1]]
        if sel in tabs:
            return [sel]
        print("  not a valid choice - try again.")


def resolve_angles(tab_name, runs, cli_angle, ask):
    """Fill the bevel angle for no-config runs that have teeth, so tooth length
    can be computed. Priority: config angle (already present) -> --angle -> prompt.
    In a non-interactive shell input() raises EOFError and the run is left blank."""
    for run in runs:
        if run.config is not None or not run.cut_teeth:
            continue
        if cli_angle is not None:
            run.manual_angle = cli_angle
            continue
        if not ask:
            continue
        try:
            s = input(f"Tab {tab_name!r} run {run.idx}: no config line - enter bevel "
                      "angle (deg from vertical) for tooth-length [blank = skip]: ").strip()
        except EOFError:
            s = ""
        if s:
            try:
                run.manual_angle = float(s)
            except ValueError:
                print("  not a number - skipping tooth-length for this run.")


def process_tab(tab_name, rows, save, outdir, cli_angle=None, ask=True):
    if not rows:
        print(f"\nTab {tab_name!r} is empty - skipping.")
        return
    runs = segment_runs(rows)
    resolve_angles(tab_name, runs, cli_angle, ask)
    summaries = print_report(tab_name, runs)
    if not save:
        return
    outdir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^0-9A-Za-z._-]", "_", tab_name)
    md = outdir / f"QA_{safe}_{stamp}.md"
    cv = outdir / f"QA_{safe}_{stamp}_teeth.csv"
    write_markdown(md, tab_name, runs, summaries)
    write_csv(cv, tab_name, runs)
    print(f"\nSaved: {md}\n       {cv}")


def main():
    ap = argparse.ArgumentParser(description="Summarise blade-grind logs from Google Sheets.")
    ap.add_argument("tab", nargs="?", help="tab name to summarise (skip the prompt)")
    ap.add_argument("--all", action="store_true", help="summarise every tab")
    ap.add_argument("--file", help="read a downloaded .xlsx instead of the live sheet")
    ap.add_argument("--id", default=SPREADSHEET_ID, help="override spreadsheet id")
    ap.add_argument("--outdir", default=str(REPORT_DIR), help="where to save reports")
    ap.add_argument("--no-save", action="store_true", help="console only, no report files")
    ap.add_argument("--angle", type=float,
                    help="bevel angle (deg from vertical) to use for no-config runs, "
                         "instead of being prompted")
    args = ap.parse_args()

    source = args.file if args.file else download_workbook(args.id)
    tabs = load_tabs(source)
    if not tabs:
        raise SystemExit("No tabs found in the workbook.")

    if args.all:
        chosen = list(tabs)
    elif args.tab:
        if args.tab not in tabs:
            raise SystemExit(f"Tab {args.tab!r} not found. Available: {', '.join(tabs)}")
        chosen = [args.tab]
    else:
        chosen = choose_tab(tabs)

    for name in chosen:
        process_tab(name, tabs[name], not args.no_save, Path(args.outdir), args.angle)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(1)
